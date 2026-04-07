#!/usr/bin/env python3
"""
Excel Engine — GitHub Copilot CLI Extension

A tool-calling extension that exposes the Excel Engine's autonomous
Excel-automation capabilities to the GitHub Copilot CLI agent.

Protocol: reads a JSON request from stdin, writes a JSON response to stdout.

Request format:
  {
    "tool": "<tool_name>",
    "params": { ... }
  }

Response format:
  {
    "status": "success" | "error",
    "result": { ... }          // on success
    "error": "<message>"       // on error
  }
"""

from __future__ import annotations

import json
import logging
import sys
import traceback
from dataclasses import asdict
from pathlib import Path
from typing import Any

# ── Ensure the excel_engine package is importable ──
# Resolve the repo root (two levels up from extensions/)
_REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO_ROOT))

from excel_engine import ExcelEngine, EngineConfig, __version__
from excel_engine.config import TaskType, Layer, TASK_LAYER_MAP
from excel_engine.parsers.instruction_parser import InstructionParser
from excel_engine.parsers.task_extractor import TaskExtractor
from excel_engine.verifier.workbook_verifier import WorkbookVerifier

# Suppress engine logs from cluttering stdout (JSON protocol)
logging.basicConfig(level=logging.WARNING, stream=sys.stderr)
logger = logging.getLogger("excel_engine_extension")

# ─────────────────────────────────────────────────────────
# Tool implementations
# ─────────────────────────────────────────────────────────


def complete_excel_assignment(params: dict[str, Any]) -> dict[str, Any]:
    """
    Run the full Excel Engine pipeline on an assignment.

    Params:
        workbook_path  (str): Path to the .xlsx workbook file.
        instruction_path (str): Path to the instruction file (.docx, .rtfd, .pdf, .txt).
    """
    workbook_path = Path(params["workbook_path"]).expanduser().resolve()
    instruction_path = Path(params["instruction_path"]).expanduser().resolve()

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    if not instruction_path.exists():
        raise FileNotFoundError(f"Instruction file not found: {instruction_path}")

    engine = ExcelEngine()
    result = engine.run(workbook=workbook_path, instructions=instruction_path)

    verifications = []
    for sv in result.verifications:
        verifications.append({
            "section_id": sv.section_id,
            "all_passed": sv.all_passed,
            "passed": sv.pass_count,
            "failed": sv.fail_count,
            "summary": sv.summary(),
        })

    return {
        "success": result.success,
        "workbook": str(result.workbook_path),
        "sections_completed": result.sections_completed,
        "sections_total": result.sections_total,
        "tasks_completed": result.tasks_completed,
        "tasks_total": result.tasks_total,
        "elapsed_seconds": round(result.elapsed_seconds, 2),
        "verifications": verifications,
        "errors": result.errors,
        "summary": result.summary(),
    }


def parse_excel_instructions(params: dict[str, Any]) -> dict[str, Any]:
    """
    Parse an instruction file and return structured tasks.

    Params:
        instruction_path (str): Path to the instruction file.
    """
    instruction_path = Path(params["instruction_path"]).expanduser().resolve()

    if not instruction_path.exists():
        raise FileNotFoundError(f"Instruction file not found: {instruction_path}")

    parser = InstructionParser()
    text = parser.parse(instruction_path)

    extractor = TaskExtractor()
    tasks = extractor.extract(text)

    task_dicts = []
    for t in tasks:
        task_dicts.append({
            "id": t.id,
            "task_type": t.task_type.value,
            "description": t.description,
            "sheet": t.sheet,
            "cell": t.cell,
            "range": t.range,
            "value": t.value,
            "formula": t.formula,
            "style": t.style,
            "params": t.params,
            "depends_on": t.depends_on,
        })

    return {
        "instruction_file": str(instruction_path),
        "raw_text_length": len(text),
        "task_count": len(tasks),
        "tasks": task_dicts,
        "task_type_summary": _summarize_task_types(tasks),
    }


def verify_excel_workbook(params: dict[str, Any]) -> dict[str, Any]:
    """
    Verify an Excel workbook's completion status.

    Params:
        workbook_path (str): Path to the .xlsx workbook file.
        instruction_path (str, optional): Path to instructions for task-level verification.
    """
    workbook_path = Path(params["workbook_path"]).expanduser().resolve()

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    verifier = WorkbookVerifier()
    verifier.load(workbook_path)

    # If instructions provided, do full task-level verification
    instruction_path = params.get("instruction_path")
    if instruction_path:
        instruction_path = Path(instruction_path).expanduser().resolve()
        if not instruction_path.exists():
            raise FileNotFoundError(f"Instruction file not found: {instruction_path}")

        parser = InstructionParser()
        text = parser.parse(instruction_path)
        extractor = TaskExtractor()
        tasks = extractor.extract(text)

        sv = verifier.verify_section("full", tasks)
        verifier.close()

        results = []
        for r in sv.results:
            results.append({
                "task_id": r.task_id,
                "task_type": r.task_type.value,
                "passed": r.passed,
                "message": r.message,
            })

        return {
            "workbook": str(workbook_path),
            "all_passed": sv.all_passed,
            "passed": sv.pass_count,
            "failed": sv.fail_count,
            "total": len(sv.results),
            "results": results,
            "summary": sv.summary(),
        }

    # Without instructions, do a structural report
    from openpyxl import load_workbook as _load_wb

    wb = _load_wb(str(workbook_path), data_only=False)
    sheets_info = []
    for name in wb.sheetnames:
        ws = wb[name]
        sheets_info.append({
            "name": name,
            "dimensions": ws.dimensions,
            "has_charts": len(ws._charts) > 0 if hasattr(ws, "_charts") else False,
            "has_tables": len(ws.tables) > 0 if hasattr(ws, "tables") else False,
            "has_merged_cells": len(ws.merged_cells.ranges) > 0,
            "has_data_validations": (
                len(ws.data_validations.dataValidation) > 0
                if ws.data_validations else False
            ),
        })
    wb.close()
    verifier.close()

    return {
        "workbook": str(workbook_path),
        "sheet_count": len(sheets_info),
        "sheets": sheets_info,
        "note": "Provide instruction_path for task-level verification.",
    }


def get_excel_engine_info(_params: dict[str, Any] | None = None) -> dict[str, Any]:
    """Return engine version, capabilities, and supported task types."""
    layers = [
        {"id": layer.value, "name": layer.name, "description": desc}
        for layer, desc in [
            (Layer.OPENPYXL, "Offline .xlsx manipulation via openpyxl"),
            (Layer.XLWINGS, "Live Excel API bridge via xlwings"),
            (Layer.APPLESCRIPT, "Excel-specific AppleScript commands"),
            (Layer.SYSTEM_EVENTS, "macOS ribbon/dialog UI automation"),
            (Layer.VBA, "VBA macro injection via VBE clipboard"),
            (Layer.PYAUTOGUI, "Last-resort desktop pixel automation"),
        ]
    ]

    task_types = sorted([tt.value for tt in TaskType])

    supported_formats = sorted(InstructionParser.SUPPORTED_EXTENSIONS)

    return {
        "version": __version__,
        "platform": "macOS",
        "requires": ["Python 3.11+", "Microsoft Excel for Mac", "macOS"],
        "layers": layers,
        "task_types": task_types,
        "task_type_count": len(task_types),
        "supported_instruction_formats": supported_formats,
        "protocol": "Engine Protocol v2.0 (SCAN → GROUP → EXECUTE → VERIFY)",
    }


# ─────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────


def _summarize_task_types(tasks) -> dict[str, int]:
    """Count tasks by type."""
    summary: dict[str, int] = {}
    for t in tasks:
        key = t.task_type.value
        summary[key] = summary.get(key, 0) + 1
    return summary


# ─────────────────────────────────────────────────────────
# Tool registry & dispatcher
# ─────────────────────────────────────────────────────────

TOOLS = {
    "complete_excel_assignment": {
        "fn": complete_excel_assignment,
        "description": (
            "Run the full Excel Engine pipeline to autonomously complete "
            "an Excel assignment. Parses instructions, plans tasks, executes "
            "across 6 automation layers, and verifies results."
        ),
        "parameters": {
            "workbook_path": {"type": "string", "required": True},
            "instruction_path": {"type": "string", "required": True},
        },
    },
    "parse_excel_instructions": {
        "fn": parse_excel_instructions,
        "description": (
            "Parse an instruction file (.docx, .rtfd, .pdf, .txt) and return "
            "a structured list of Excel tasks with types and dependencies."
        ),
        "parameters": {
            "instruction_path": {"type": "string", "required": True},
        },
    },
    "verify_excel_workbook": {
        "fn": verify_excel_workbook,
        "description": (
            "Verify an Excel workbook's completion status. Optionally provide "
            "instruction_path for full task-level pass/fail verification."
        ),
        "parameters": {
            "workbook_path": {"type": "string", "required": True},
            "instruction_path": {"type": "string", "required": False},
        },
    },
    "get_excel_engine_info": {
        "fn": get_excel_engine_info,
        "description": (
            "Return engine version, supported task types, automation layers, "
            "and capability information."
        ),
        "parameters": {},
    },
}


def list_tools() -> dict[str, Any]:
    """Return the tool manifest (without function references)."""
    manifest = {}
    for name, spec in TOOLS.items():
        manifest[name] = {
            "description": spec["description"],
            "parameters": spec["parameters"],
        }
    return {"tools": manifest}


def dispatch(request: dict[str, Any]) -> dict[str, Any]:
    """Route a request to the appropriate tool function."""
    tool_name = request.get("tool")

    if tool_name == "list_tools":
        return {"status": "success", "result": list_tools()}

    if tool_name not in TOOLS:
        available = ", ".join(sorted(TOOLS.keys()))
        return {
            "status": "error",
            "error": f"Unknown tool: {tool_name!r}. Available: {available}",
        }

    params = request.get("params", {})
    tool_fn = TOOLS[tool_name]["fn"]
    result = tool_fn(params)
    return {"status": "success", "result": result}


# ─────────────────────────────────────────────────────────
# Main — stdin/stdout JSON protocol
# ─────────────────────────────────────────────────────────


def main() -> None:
    """Read JSON request from stdin, write JSON response to stdout."""
    raw = sys.stdin.read().strip()

    if not raw:
        # No input — print tool manifest for discovery
        response = {"status": "success", "result": list_tools()}
        print(json.dumps(response, indent=2))
        return

    try:
        request = json.loads(raw)
    except json.JSONDecodeError as e:
        response = {"status": "error", "error": f"Invalid JSON input: {e}"}
        print(json.dumps(response))
        sys.exit(1)

    try:
        response = dispatch(request)
    except FileNotFoundError as e:
        response = {"status": "error", "error": str(e)}
    except Exception as e:
        logger.error("Tool execution failed: %s", e, exc_info=True)
        response = {
            "status": "error",
            "error": str(e),
            "traceback": traceback.format_exc(),
        }

    print(json.dumps(response, indent=2, default=str))


if __name__ == "__main__":
    main()
